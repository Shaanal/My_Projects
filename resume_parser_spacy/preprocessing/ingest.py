import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from PIL import Image

IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
SUPPORTED = IMAGES | {".pdf", ".docx", ".txt", ".md", ".csv", ".tsv", ".xml", ".json", ".html", ".htm"}

GUTTER_TOLERANCE = 12 
LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        if data.strip():
            self.parts.append(data.strip())


# ---------- PDF ----------

def _crosses(block, x):
    """True for a full-width block (e.g. the name banner), not a column block that slightly overhangs x."""
    return block[0] < x - GUTTER_TOLERANCE and block[2] > x + GUTTER_TOLERANCE


def _column_split(blocks, width):
    """x position of the gap between two text columns, or None for a one-column page."""
    total = sum(len(b[4]) for b in blocks) or 1
    scored = []
    for x in range(int(width * 0.25), int(width * 0.75), 2):
        crossing = sum(len(b[4]) for b in blocks if b[0] < x < b[2])
        left = sum(len(b[4]) for b in blocks if b[2] <= x)
        right = sum(len(b[4]) for b in blocks if b[0] >= x)
        if crossing <= 0.15 * total and min(left, right) >= 0.15 * total:
            scored.append((crossing, x))
    if not scored:
        return None

    # Middle of the widest band of positions crossed by the least text.
    least = min(c for c, _ in scored)
    xs = [x for c, x in scored if c == least]
    runs = [[xs[0]]]
    for x in xs[1:]:
        if x - runs[-1][-1] == 2:
            runs[-1].append(x)
        else:
            runs.append([x])
    widest = max(runs, key=len)
    return widest[len(widest) // 2]


def _page_text(page):
    blocks = [b for b in page.get_text("blocks", sort=True) if b[6] == 0 and b[4].strip()]
    if not blocks:
        return ""
    split = _column_split(blocks, page.rect.width)
    if split is None:
        return "\n".join(b[4].strip() for b in blocks)

    # Between two full-width blocks, read the left column fully, then the right one.
    ordered, segment = [], []

    def flush():
        by_position = lambda b: (b[1], b[0])
        ordered.extend(sorted((b for b in segment if (b[0] + b[2]) / 2 <= split), key=by_position))
        ordered.extend(sorted((b for b in segment if (b[0] + b[2]) / 2 > split), key=by_position))
        segment.clear()

    for block in sorted(blocks, key=lambda b: (b[1], b[0])):
        if _crosses(block, split):
            flush()
            ordered.append(block)
        else:
            segment.append(block)
    flush()
    return "\n".join(b[4].strip() for b in ordered)


def _pdf_text(path):
    pages = []
    for page in pymupdf.open(str(path)):
        text = _page_text(page)
        if len(text.strip()) < 40:  # scanned page: OCR it
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            text = pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))))
        pages.append(text)
    return "\n".join(pages).strip()


# ---------- DOCX ----------

def _paragraph_lines(element):
    """Every paragraph under element, including text boxes (skipping Word's duplicate fallback copies)."""
    lines = []
    for para in element.iter(W + "p"):
        if any(a.tag.endswith("}Fallback") for a in para.iterancestors()):
            continue
        parts = []
        for node in para.iter(W + "t", W + "tab", W + "br", W + "cr"):
            if next(node.iterancestors(W + "p")) is not para:
                continue
            if node.tag == W + "t":
                parts.append(node.text or "")
            else:
                parts.append(" " if node.tag == W + "tab" else "\n")
        lines.extend(x.strip() for x in "".join(parts).splitlines() if x.strip())
    return lines


def _docx_text(path):
    doc = Document(str(path))
    lines = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag in {"p", "sdt"}:
            lines.extend(_paragraph_lines(child))
        elif tag == "tbl":
            table = next(t for t in doc.tables if t._tbl is child)
            for row in table.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    if id(cell._tc) not in seen:  # merged cells repeat
                        seen.add(id(cell._tc))
                        cells.append(cell.text.strip())
                # Multi-line cells are layout columns: read each one top to bottom.
                if any("\n" in c for c in cells):
                    lines.extend(c for c in cells if c)
                elif any(cells):
                    lines.append(" | ".join(c for c in cells if c))
    return "\n".join(lines)


# ---------- public functions ----------

def extract_text(file_path: str) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {', '.join(sorted(SUPPORTED))}")

    if ext == ".pdf":
        return _pdf_text(path)
    if ext == ".docx":
        return _docx_text(path)
    if ext in IMAGES:
        return pytesseract.image_to_string(Image.open(str(path)))

    raw = path.read_text(encoding="utf-8", errors="ignore")
    if ext in {".html", ".htm", ".xml"}:
        parser = _HTMLText()
        parser.feed(raw)
        return "\n".join(parser.parts)
    if ext == ".json":
        try:
            return json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            return raw
    return raw


def extract_links(file_path: str) -> list:
    """Hyperlinks stored in the file (often hidden behind icons or a "LinkedIn" label)."""
    path = Path(file_path)
    ext = path.suffix.lower()
    links = []
    try:
        if ext == ".pdf":
            for page in pymupdf.open(str(path)):
                links += [link["uri"] for link in page.get_links() if link.get("uri")]
        elif ext == ".docx":
            rels = Document(str(path)).part.rels.values()
            links += [r.target_ref for r in rels if r.reltype.endswith("/hyperlink") and r.is_external]
        elif ext in {".html", ".htm"}:
            html = path.read_text(encoding="utf-8", errors="ignore")
            links += re.findall(r"""href=["']([^"']+)["']""", html)
    except Exception:
        return []
    return list(dict.fromkeys(x.strip() for x in links if x.strip()))


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    for ligature, letters in LIGATURES.items():
        text = text.replace(ligature, letters)
    text = text.replace(" ", " ").replace("", "•")

    # Icon-font glyphs (calendar, phone, pin) arrive as control or private-use characters.
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f-]", " ", text)
    # OCR often reads a round bullet as "e", "¢", "©" or "«".
    text = re.sub(r"(?m)^[ \t]*(?:e|¢|©|«|°)[ \t]+(?=[A-Z])", "• ", text)

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(?<=[a-z])-\n(?=[a-z])", "", text)  # "devel-\nopment" -> "development"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def restore_line_breaks(text: str) -> str:
    """The XML corpus stores line breaks as runs of spaces; turn the last space into "\\n" (offsets unchanged)."""
    return re.sub(r" {2,}", lambda m: " " * (len(m.group()) - 1) + "\n", text)
